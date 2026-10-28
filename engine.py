import re,io,os,json,sqlite3,datetime,hashlib
import fitz
from docx import Document
from openpyxl import load_workbook,Workbook
from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,Table,TableStyle,PageBreak
from reportlab.lib.styles import getSampleStyleSheet,ParagraphStyle
from reportlab.lib import colors
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from xml.sax.saxutils import escape
SECTIONS={
'1':('前言',['公司介紹與營運概況','報導個體與揭露範圍','IFRS 永續揭露準則導入情形','遵循聲明與編製基礎','專章核准程序與適用豁免']),
'2':('治理',['永續及氣候相關治理架構','治理單位之監督職責','管理階層之執行職責','永續相關專業能力','永續相關決策與政策權衡','永續目標、績效考核與薪酬連結']),
'3':('策略',['永續及氣候相關風險與機會','短、中、長期時間區間','對經營模式與價值鏈之影響','永續及氣候相關因應策略','永續相關風險與機會之財務影響','氣候情境分析與策略韌性']),
'4':('風險管理',['永續相關風險管理架構','風險與機會辨認流程','風險評估與排序','氣候相關風險管理','其他永續相關風險管理','風險監控與韌性管理']),
'5':('指標與目標',['永續相關指標與目標總覽','溫室氣體排放指標','氣候相關跨行業指標','行業基礎指標','其他永續績效指標','目標達成情形與未來規劃'])}
RULES={
'1.1':'公司簡介 公司介紹 公司概況 主要業務 產品服務 營運模式 價值鏈 公司沿革',
'1.2':'合併財務報告 報導範圍 子公司 合併範圍 報導期間 組織邊界 統計邊界',
'1.3':'IFRS S1 S2 導入計畫 永續資訊內控 導入時程',
'1.4':'遵循聲明 編製基礎 衡量基礎 重大估計 編製依據',
'1.5':'董事會通過 核准日期 核准程序 豁免 過渡規定',
'2.1':'治理架構 組織圖 永續委員會 董事會 功能性委員會 權責分工',
'2.2':'董事會議紀錄 監督職責 呈報頻率 治理單位 監督機制',
'2.3':'管理階層 執行委員會 跨部門 內部控制 管理職責',
'2.4':'董事專業能力 教育訓練 進修時數 永續課程 專業技能',
'2.5':'決策權衡 政策權衡 重大決策 風險與機會權衡',
'2.6':'薪酬連結 績效考核 獎酬 績效指標 永續目標',
'3.1':'氣候風險 永續風險 氣候機會 永續機會 重大風險',
'3.2':'短期 中期 長期 時間區間 規劃週期',
'3.3':'上游 下游 價值鏈 營運影響 供應鏈影響',
'3.4':'淨零轉型 減碳措施 能源管理 轉型計畫 數位轉型',
'3.5':'財務影響 現金流量 資金成本 籌資 財務績效 資產負債',
'3.6':'情境分析 SSP RCP NGFS 氣候韌性 實體風險 轉型風險',
'4.1':'風險管理制度 風險管理政策 風險管理架構 ERM',
'4.2':'風險辨認 風險識別 評估方法 資料來源 參數假設',
'4.3':'可能性 影響程度 風險排序 風險矩陣 財務衝擊',
'4.4':'實體風險 轉型風險 氣候風險管理',
'4.5':'資訊安全 法規遵循 人力資源 營運中斷 聲譽風險',
'4.6':'風險監控 監控頻率 改善措施 風險追蹤 韌性管理',
'5.1':'ESG目標 目標追蹤 基準年 達成率 執行成果 管理指標',
'5.2':'溫室氣體 範疇一 範疇二 範疇三 scope 1 scope 2 scope 3 tCO2e 碳盤查 排放量',
'5.3':'內部碳價 資本配置 跨行業指標 氣候相關機會 氣候轉型風險',
'5.4':'SASB 行業基礎指標 產業指標 行業揭露',
'5.5':'用水量 廢棄物 員工人數 職業安全 資安事件 能源使用 供應鏈',
'5.6':'目標達成 目標差異 未來規劃 改善計畫 減量目標'}
KEYS={k:re.split(r'\s+',v.lower()) for k,v in RULES.items()}
LABELS={'E':'環境','S':'社會','G':'治理'}
def classify(text,name=''):
 t=(text+' '+name).lower(); scores={k:sum(min(t.count(w),4)*(3 if w in name.lower() else 1) for w in words) for k,words in KEYS.items()}
 ranked=sorted(scores,key=lambda k:scores[k],reverse=True)
 best=scores[ranked[0]]
 if not best:return [],0
 chosen=[k for k in ranked if scores[k]>=max(1,best*.58)][:4]
 return chosen,round(min(.95,.35+best/(best+12)),2)
def category(t):
 return 'E' if re.search('排放|能源|氣候|減碳|水資源|廢棄物',t) else ('S' if re.search('員工|職安|人才|社會|供應商',t) else 'G')
def extract(filename,data):
 ext=filename.rsplit('.',1)[-1].lower(); chunks=[]
 if ext=='pdf':
  d=fitz.open(stream=data,filetype='pdf')
  for i,p in enumerate(d):
   t=p.get_text(sort=True).strip()
   if t:chunks.append((f'第 {i+1} 頁',t))
 elif ext=='docx':
  d=Document(io.BytesIO(data)); chunks=[('內文', '\n'.join(x.text for x in d.paragraphs if x.text.strip()))]
  for i,tab in enumerate(d.tables):chunks.append((f'表格 {i+1}','\n'.join(' | '.join(c.text for c in r.cells) for r in tab.rows)))
 elif ext in ('xlsx','xlsm'):
  w=load_workbook(io.BytesIO(data),read_only=True,data_only=True)
  for s in w:
   rows=[' | '.join(str(v) if v is not None else '' for v in row) for row in s.iter_rows(values_only=True)]
   for i in range(0,len(rows),35):chunks.append((f'{s.title} 第 {i+1}-{min(i+35,len(rows))} 列','\n'.join(rows[i:i+35])))
 elif ext in ('txt','csv','md'):
  t=data.decode('utf-8-sig',errors='replace');chunks=[('全文',t)]
 else:raise ValueError('不支援的格式；掃描圖片需另接 OCR')
 result=[]
 for loc,t in chunks:
  for i in range(0,len(t),2600):
   piece=t[i:i+2600].strip()
   if piece:result.append((loc+(f' 片段 {i//2600+1}' if len(t)>2600 else ''),piece))
 return result

def db(path):
 c=sqlite3.connect(path);c.row_factory=sqlite3.Row
 c.executescript('''CREATE TABLE IF NOT EXISTS projects(id INTEGER PRIMARY KEY,name TEXT,industry TEXT,year TEXT,created TEXT);
 CREATE TABLE IF NOT EXISTS sources(id INTEGER PRIMARY KEY,project INTEGER,name TEXT,sha TEXT,created TEXT);
 CREATE TABLE IF NOT EXISTS evidence(id INTEGER PRIMARY KEY,project INTEGER,source INTEGER,location TEXT,body TEXT,sections TEXT,confidence REAL,category TEXT);
 CREATE TABLE IF NOT EXISTS drafts(project INTEGER,section TEXT,body TEXT,confirmed INTEGER DEFAULT 0,updated TEXT,PRIMARY KEY(project,section));''');return c

def generate(c,pid):
 evidence=c.execute('SELECT e.*,s.name filename FROM evidence e JOIN sources s ON s.id=e.source WHERE e.project=?',(pid,)).fetchall();made=0
 for chapter,(heading,subs) in SECTIONS.items():
  for i,title in enumerate(subs,1):
   key=f'{chapter}.{i}'; existing=c.execute('SELECT confirmed FROM drafts WHERE project=? AND section=?',(pid,key)).fetchone()
   if existing and existing['confirmed']:continue
   relevant=[r for r in evidence if key in json.loads(r['sections'])]
   if relevant:
    lines=[]
    for r in relevant[:5]:
     # source-backed excerpts, explicitly NOT fabricated prose
     compact=' '.join(r['body'].split());excerpt=compact[:420]
     lines.append(f'【待人工確認之來源摘錄】{excerpt}{"…" if len(compact)>420 else ""}\n來源：{r["filename"]}／{r["location"]}（證據 #{r["id"]}）')
    body='\n\n'.join(lines)+'\n\n[待確認：上述資料是否完整、適用及可作正式揭露；目前為來源摘錄，非 AI 撰寫之定稿]'
   else:body=f'[待補充：{title}之企業佐證資料及適用性判斷]'
   c.execute('INSERT INTO drafts(project,section,body,confirmed,updated) VALUES(?,?,?,0,?) ON CONFLICT(project,section) DO UPDATE SET body=excluded.body,confirmed=0,updated=excluded.updated',(pid,key,body,datetime.datetime.now().isoformat()));made+=1
 c.commit();return made

def report_rows(c,pid):
 rows=c.execute('SELECT * FROM drafts WHERE project=?',(pid,)).fetchall();return {r['section']:dict(r) for r in rows}
def docx_bytes(c,pid):
 p=c.execute('SELECT * FROM projects WHERE id=?',(pid,)).fetchone();r=report_rows(c,pid);d=Document();d.add_heading(f'{p["name"]} 永續相關財務資訊專章',0);d.add_paragraph(f'報導年度：{p["year"]}｜草稿，未經準則符合性確信')
 for ch,(title,subs) in SECTIONS.items():
  d.add_heading(f'{ch} {title}',1)
  for i,sub in enumerate(subs,1):
   k=f'{ch}.{i}';d.add_heading(f'{k} {sub}',2)
   for para in r.get(k,{}).get('body','[待補充]').split('\n'):d.add_paragraph(para)
 d.add_heading('附錄｜資料來源與分類紀錄',1)
 for e in c.execute('SELECT e.*,s.name filename FROM evidence e JOIN sources s ON e.source=s.id WHERE e.project=?',(pid,)):
  d.add_paragraph(f'#{e["id"]} {e["filename"]} / {e["location"]} → {e["sections"]}')
 out=io.BytesIO();d.save(out);return out.getvalue()
def excel_bytes(c,pid):
 w=Workbook();s=w.active;s.title='待補件與檢核';s.append(['章節','項目','企業資料支持','已生成文字','人工確認','待補充類型','IFRS準則符合性'])
 drafts=report_rows(c,pid)
 for ch,(title,subs) in SECTIONS.items():
  for i,sub in enumerate(subs,1):
   k=f'{ch}.{i}';n=c.execute('SELECT COUNT(*) FROM evidence WHERE project=? AND sections LIKE ?',(pid,f'%"{k}"%')).fetchone()[0];r=drafts.get(k,{});s.append([k,sub,'有' if n else '無','有' if r and not r['body'].startswith('[待補充') else '無','已確認' if r.get('confirmed') else '待審核','管理程序/數據待判斷' if not n else '來源適用性待核','未評估'])
 s2=w.create_sheet('來源追溯');s2.append(['證據ID','檔名','位置','章節','ESG標籤','分類信心','來源內容'])
 for e in c.execute('SELECT e.*,s.name filename FROM evidence e JOIN sources s ON e.source=s.id WHERE e.project=?',(pid,)):
  s2.append([e['id'],e['filename'],e['location'],e['sections'],e['category'],e['confidence'],e['body'][:30000]])
 for sheet in w:
  sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
  from openpyxl.styles import Font,PatternFill
  for cell in sheet[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='146B63')
  for col in 'ABCDEFG':sheet.column_dimensions[col].width=22
 out=io.BytesIO();w.save(out);return out.getvalue()
